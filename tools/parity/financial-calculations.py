#!/usr/bin/env python3
"""Editable-workbook regression checks, independent of screen/Excel parity.

Run after `node tools/parity/gen-xlsx.js office`:
    python3 tools/parity/financial-calculations.py

The generated workbook is never saved or rewritten. The formulas evaluator changes
assumption inputs in memory, as a user would in Excel. Checks use cash conservation,
loan payoff identities and a separate full-model scenario for each sensitivity cell;
they do not reuse the application's engine results or the generator's expressions.
"""
import argparse
import json
import math
import pathlib
import re
import sys

import formulas
import openpyxl
from workbook_layout import WorkbookLayout


ROOT = pathlib.Path(__file__).resolve().parent
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--xlsx", type=pathlib.Path, default=ROOT / "out" / "office_parity.xlsx")
parser.add_argument("--report", type=pathlib.Path, help="Optional JSON evidence file")
args = parser.parse_args()
if not args.xlsx.is_file():
    parser.error("Generate office_parity.xlsx first: node tools/parity/gen-xlsx.js office")

workbook = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=False)
layout = WorkbookLayout(workbook.sheetnames)
workbook.close()
model = formulas.ExcelModel().loads(str(args.xlsx.resolve())).finish()
base = model.calculate()
keys = {}
for key in base:
    match = re.search(r"\]([^']+)'!([A-Z]+[0-9]+)$", key)
    if match:
        keys[match[1].upper() + "!" + match[2]] = key

checks = []
scenarios = []
ERRORS = ("#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#NULL!")


def raw(sol, sheet, ref):
    key = keys.get(layout.key(sheet, ref))
    return sol[key].value[0][0] if key in sol else None


def num(sol, sheet, ref):
    try:
        return float(raw(sol, sheet, ref))
    except (TypeError, ValueError):
        # Preserve the rest of the evidence when a regression produces #NUM! or
        # a missing value. Its scenario error scan and numeric checks both fail.
        return math.nan


def check(name, passed, detail=""):
    checks.append({"name": name, "passed": bool(passed), "detail": detail})
    print(("  PASS " if passed else "  FAIL ") + name + (": " + detail if detail else ""))


def close(name, actual, expected, tolerance=1e-7):
    passed = math.isfinite(actual) and math.isfinite(expected) and abs(actual - expected) <= tolerance
    check(name, passed, "actual=%.10f expected=%.10f" % (actual, expected))


def calculate(name, overrides):
    # Every call receives only its own changes. No preceding scenario carries over.
    def input_key(ref):
        sheet, cell = ref.split("!", 1) if "!" in ref else ("01_Assumptions", ref)
        return keys[layout.key(sheet, cell)]
    sol = model.calculate(inputs={input_key(ref): value for ref, value in overrides.items()})
    scenarios.append({"name": name, "inputs": overrides})
    errors = []
    for key, value in sol.items():
        text = str(value.value)
        if any(token in text for token in ERRORS):
            errors.append(key)
    check(name + " / no Excel errors", not errors, ", ".join(errors[:5]))
    return sol


HOLD = int(num(base, "01_Assumptions", "C79"))
if HOLD != 5:
    parser.error("These regression cases require the five-year office example")
YEARS = "CDEFG"

# A three-year loan must be fully repaid in year 3 even in a five-year model.
# This catches negative outstanding balances/negative interest that still made
# the old validation sheet report PASS under low leverage.
print("\n[1] Loan payoff and payment identities")
debt_cases = [
    ("straight principal", "원금균등", 3, 0, 0.042),
    ("level payment", "원리금균등", 3, 0, 0.042),
    ("grace then level payment", "거치후 원리금균등", 3, 1, 0.042),
    ("grace equals term", "거치후 원리금균등", 3, 3, 0.042),
    ("grace exceeds term", "거치후 원리금균등", 3, 4, 0.042),
    ("zero interest", "원리금균등", 3, 0, 0),
    ("minimum term", "원리금균등", 0, 0, 0.042),
]
for name, method, term, grace, rate in debt_cases:
    s = calculate(name, {"C45": 0.1, "C47": method, "C48": term, "C80": grace, "C46": rate})
    loan = num(s, "01_Assumptions", "C49")
    principal = [num(s, "05_Debt_Schedule", col + "7") for col in YEARS]
    interest = [num(s, "05_Debt_Schedule", col + "6") for col in YEARS]
    ending = [num(s, "05_Debt_Schedule", col + "8") for col in YEARS]
    close(name + " / lifetime principal equals initial loan", math.fsum(principal), loan)
    check(name + " / no negative principal, interest or balance", all(math.isfinite(x) and x >= -1e-8 for x in principal + interest + ending))
    maturity = max(1, term)
    check(name + " / balance zero from maturity", all(abs(x) < 1e-7 for x in ending[maturity - 1:]))
    check(name + " / no payments after payoff", all(abs(x) < 1e-7 for x in principal[maturity:] + interest[maturity:]))
    if method == "원금균등":
        check(name + " / equal principal installments", all(abs(x - loan / term) < 1e-7 for x in principal[:term]))
    elif term == 3 and grace == 0:
        payments = [p + i for p, i in zip(principal[:term], interest[:term])]
        check(name + " / equal total installments", max(payments) - min(payments) < 1e-7)
    elif method == "거치후 원리금균등" and grace >= term:
        check(name + " / oversized grace ends before maturity", all(abs(x) < 1e-7 for x in principal[:term - 1]))
        close(name + " / maturity pays remaining loan", principal[term - 1], loan)

# If there is only common equity, every common cash flow and its IRR must equal
# total equity. Negative operating cash needs belong to common funding calls;
# they must never manufacture preferred distributions or preferred arrears.
print("\n[2] Waterfall conservation and zero-preferred identity")
for cumulative in (0, 1):
    name = "no preferred, negative operating cash, cumulative=%d" % cumulative
    s = calculate(name, {"C51": 0, "C73": cumulative, "C77": 4000})
    pref_cf = [num(s, "08_Equity_Cashflow", col + "5") for col in "CDEFGH"]
    common_cf = [num(s, "08_Equity_Cashflow", col + "6") for col in "CDEFGH"]
    total_cf = [num(s, "08_Equity_Cashflow", col + "7") for col in "CDEFGH"]
    check(name + " / every preferred cash flow is zero", all(abs(x) < 1e-8 for x in pref_cf))
    check(name + " / no nonexistent preferred arrears", all(abs(num(s, "07_Equity_Waterfall", col + "10")) < 1e-8 for col in YEARS))
    check(name + " / common CF equals total CF in every year", all(abs(a - b) < 1e-7 for a, b in zip(common_cf, total_cf)))
    check(name + " / actual common funding call exists", common_cf[1] < -1)
    close(name + " / common IRR equals total IRR", num(s, "09_Return_Summary", "C5"), num(s, "09_Return_Summary", "E5"), 1e-9)
    for label, flows, ref in (("common", common_cf, "C7"), ("total", total_cf, "E7")):
        # Cash-flow row includes the initial investment. Later negative cash
        # flows are further contributions, not reductions of gross receipts.
        receipts = math.fsum(amount for amount in flows if amount > 0)
        contributions = math.fsum(-amount for amount in flows if amount < 0)
        close(name + " / " + label + " EM includes all funding calls",
              num(s, "09_Return_Summary", ref), receipts / contributions, 1e-9)

for cumulative in (0, 1):
    name = "preferred present, negative operating cash, cumulative=%d" % cumulative
    s = calculate(name, {"C73": cumulative, "C77": 4000})
    check(name + " / preferred payments are never negative", all(num(s, "07_Equity_Waterfall", col + "8") >= -1e-8 for col in YEARS))
    check(name + " / distributions conserve available cash", all(abs(num(s, "07_Equity_Waterfall", col + "8") + num(s, "07_Equity_Waterfall", col + "9") - num(s, "07_Equity_Waterfall", col + "5")) < 1e-7 for col in YEARS))
    coupon = num(s, "01_Assumptions", "C53") * num(s, "01_Assumptions", "C52")
    close(name + " / first loss-year arrears contain only unpaid coupon", num(s, "07_Equity_Waterfall", "C10"), coupon if cumulative else 0)

# The sensitivity table is compared with a fresh run of the main workbook using
# the actual row/column assumption, not with an alternate copy of its formula.
# Fees and rent-linked deposits expose errors hidden by the default zero fee.
print("\n[3] Sensitivity scenarios versus full-model recalculation")
fee = {"C72": 0.2}
table = calculate("sensitivity with 20% performance fee", fee)
base_rent = num(table, "01_Assumptions", "C23")
base_cap = num(table, "01_Assumptions", "C58")
close("sensitivity centre equals main IRR", num(table, "10_Sensitivity", "E15"), num(table, "09_Return_Summary", "E5"), 1e-8)
for label, table_ref, rent_scale, cap_ref in [
    ("cap decrease with fee", "C7", 1, "C5"),
    ("cap increase with fee", "G7", 1, "G5"),
    ("rent increase including deposits", "E17", 1.1, "E5"),
    ("rent decrease plus cap increase", "G13", 0.9, "G5"),
    ("rent increase plus cap decrease", "C17", 1.1, "C5"),
]:
    target_cap = num(table, "10_Sensitivity", cap_ref)
    s = calculate(label, {**fee, "C23": base_rent * rent_scale, "C58": target_cap})
    close(label + " / table IRR equals main model", num(table, "10_Sensitivity", table_ref), num(s, "09_Return_Summary", "E5"), 1e-8)
    if rent_scale != 1:
        check(label + " / opening deposit actually changes", abs(num(s, "01_Assumptions", "C75") - num(table, "01_Assumptions", "C75")) > 1)
        check(label + " / initial equity actually changes", abs(num(s, "01_Assumptions", "C55") - num(table, "01_Assumptions", "C55")) > 1)
    if table_ref == "C7":
        # Also verify the companion EM and sale-value rows use that full scenario.
        close(label + " / table EM equals main model", num(table, "10_Sensitivity", "C8"), num(s, "09_Return_Summary", "E7"), 1e-8)
        close(label + " / table sale value equals main model", num(table, "10_Sensitivity", "C6"), num(s, "06_Tax_Disposition", "C13"), 1e-6)

# Class-specific after-tax distributions are not implemented. Blank is honest;
# copying before-tax class IRRs under a 'after-tax' heading is not.
print("\n[4] Tax calculation scope and distribution policy")
for passthrough in (0, 1):
    name = "tax mode %d" % passthrough
    s = calculate(name, {"C66": passthrough})
    for ref in ("C6", "D6"):
        value = raw(s, "09_Return_Summary", ref)
        check(name + " / unsupported class after-tax " + ref + " is blank", value is None or str(value).strip() == "")
    pretax, aftertax = (num(s, "09_Return_Summary", ref) for ref in ("E5", "E6"))
    if passthrough:
        close(name + " / total-equity tax exemption", aftertax, pretax, 1e-9)
    else:
        check(name + " / total-equity corporate tax has an effect", aftertax < pretax - 0.001)

workbook = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=False)
assumptions = workbook["A&R" if layout.compact else "01_Assumptions"]
check("unused payout percentage is replaced with distribution policy", assumptions["C67"].value == "전액 분배")
if layout.compact:
    policy = " ".join(str(assumptions[ref].value) for ref in ("C67", "E67"))
    check("distribution policy explains distribution and common funding", all(text in policy for text in ("전액 분배", "보통주", "추가출자")))
else:
    check("distribution policy is identified as calculation basis", assumptions["E67"].value == "계산 기준")
check("distribution policy is no longer styled as an input", assumptions["C67"].font.color != assumptions["C65"].font.color)
workbook.close()

# A validator must fail when its invariant is deliberately violated. Overriding
# calculated values exists only inside the evaluator and never modifies the XLSX.
# Testing just an undisturbed PASS would not prove detection or top-level wiring.
print("\n[5] Validation checks detect injected financial inconsistencies")
for ref in ("E27", "E28", "E29", "E30", "E15"):
    check("valid baseline / " + ref + " is PASS", str(raw(base, "11_Validation_Checks", ref)).strip() == "PASS")
for name, overrides, verdict in [
    ("negative ending debt", {"05_Debt_Schedule!G8": -1}, "E27"),
    ("common cash flow differs from total", {"08_Equity_Cashflow!D6": num(base, "08_Equity_Cashflow", "D6") + 1}, "E28"),
    ("preferred payment with no preferred capital", {"C51": 0, "08_Equity_Cashflow!D5": -1}, "E29"),
    ("unsupported repayment label", {"C47": "만기일시"}, "E30"),
    ("invalid preferred cumulative flag", {"C73": 2}, "E30"),
    ("negative preferred capital", {"C51": -0.1}, "E13"),
]:
    s = calculate(name, overrides)
    check(name + " / dedicated check fails", str(raw(s, "11_Validation_Checks", verdict)).strip() == "FAIL")
    check(name + " / overall verdict fails", str(raw(s, "11_Validation_Checks", "E15")).strip() == "FAIL")

failed = sum(not item["passed"] for item in checks)
report = {"xlsx": str(args.xlsx.resolve()), "scenarios": scenarios, "checks": checks,
          "passed": len(checks) - failed, "failed": failed}
if args.report:
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print("\nFINANCIAL CALCULATIONS %s - %d checks, %d scenarios, %d failures" %
      ("FAIL" if failed else "OK", len(checks), len(scenarios), failed))
sys.exit(1 if failed else 0)
